"""全国の自治体台帳を作る

  python nationwide/build_registry.py
  python nationwide/build_registry.py --no-jlis-fallback    # J-LIS の URL を 2 つ目の公式サイトに足さない (前の動き)
  (先に python nationwide/fetch_registry_sources.py と python nationwide/islands.py)

出力:
  out/registry/municipalities.csv   1 行 = 1 自治体コード (1,741 + 北方領土 6 村)
  out/registry/municipalities.json  同じ内容を入れ子のまま + 出典と件数
  out/registry/no_service_areas.json  gps2city がコードを返しても gomisute-gps が答えない所
  out/registry/designated_islands_n03.csv  指定離島と N03 の陸のかたまりの対応 (当てられないものは空)

入力と出典は nationwide/README.md。件数が調査の数と合わなければ止まる (assert)。
"""
import sys, re, io, csv, json, html, glob, math, pathlib, zipfile, datetime, itertools, collections, urllib.parse, warnings
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
REG = ROOT / "data" / "registry"
WORK = REG / "work"
OUT = ROOT / "out" / "registry"
HERE = ROOT / "nationwide"

NT = {"016951", "016969", "016977", "016985", "016993", "017001"}   # 北方領土の 6 村
DATES = {"172014": "金沢市", "045811": "女川町"}                      # gomisute-gps が日付で答えられる所
SHOZOKU_MITEI = {"12000", "13000", "23000", "30000", "40000", "46000", "47000"}
J_LANG = "Q5287"                                                     # 日本語 (P407 の値)


def check_digit(code5):
    s = sum(int(c) * w for c, w in zip(code5, (6, 5, 4, 3, 2)))
    return code5 + str((11 - s % 11) % 10)


def zip_csv_rows(path):
    def walk(b):
        with zipfile.ZipFile(io.BytesIO(b)) as z:
            for n in z.namelist():
                d = z.read(n)
                if n.endswith(".zip"):
                    yield from walk(d)
                elif n.endswith(".csv"):
                    yield d
    for d in walk(pathlib.Path(path).read_bytes()):
        yield from csv.DictReader(io.StringIO(d.decode("utf-8-sig")))


def read_csv(path):
    return list(csv.DictReader(open(path, encoding="utf-8")))


# ---------- 総務省の団体コード (台帳の行) ----------
def load_soumu():
    wb = openpyxl.load_workbook(REG / "soumu" / "000925835.xlsx", read_only=True)
    munis = []
    for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
        if r and r[0] and r[2]:
            code, pref, name = str(r[0]).zfill(6), r[1], r[2]
            if code in NT:
                kind = "北方領土"
            elif pref == "東京都" and name.endswith("区"):
                kind = "特別区"
            else:
                kind = name[-1]
            munis.append({"lg_code": code, "pref": pref, "name": name, "kind": kind})
    codes = {m["lg_code"] for m in munis}
    wards, city = collections.defaultdict(list), None
    for r in wb.worksheets[1].iter_rows(min_row=2, values_only=True):
        if not (r and r[0]):
            continue
        code = str(r[0]).zfill(6)
        if code in codes:
            city = code
        else:
            wards[city].append(code)
    return munis, dict(wards)


# ---------- Wikidata P856 と J-LIS ----------
def norm_url(u):
    """公式サイトの URL をそろえる (小文字のホスト、既定のポートを外す、パスが空なら /、# 以降を外す)"""
    p = urllib.parse.urlsplit(u.strip())
    scheme = (p.scheme or "http").lower()
    host = (p.hostname or "").lower()
    port = p.port
    netloc = host + (f":{port}" if port and (scheme, port) not in (("http", 80), ("https", 443)) else "")
    return urllib.parse.urlunsplit((scheme, netloc, p.path or "/", p.query, ""))


def url_key(u):
    p = urllib.parse.urlsplit(u)
    host = (p.hostname or "").removeprefix("www.")
    path = re.sub(r"/index\.(html?|php)$", "/", p.path).rstrip("/").lower()
    return host, path


def host_of(u):
    return (urllib.parse.urlsplit(u).hostname or "").removeprefix("www.")


GOV_HOST = re.compile(r"\.(lg|go)\.jp$|(^|\.)(city|town|vill|village|pref)\.")


def gov_like(u):
    """自治体のホストらしいか (lg.jp、または city. / town. / vill. を含む)。
    これに当たらない公式サイトも多い (北海道の町など) ので、これだけで外すことはしない"""
    return bool(GOV_HOST.search(host_of(u)))


def order_sites(urls):
    """同じ所を指す URL を 1 つにし (https を残す)、lg.jp を先に並べる"""
    best = {}
    for u in urls:
        k = url_key(u)
        if k not in best or (u.startswith("https:") and not best[k].startswith("https:")):
            best[k] = u
    return sorted(best.values(), key=lambda u: (not host_of(u).endswith(".lg.jp"), not u.startswith("https:"),
                                                len(urllib.parse.urlsplit(u).path), u))


def load_wikidata():
    """wikidata.org から取った今の項目 (entities.jsonl) の P429 と P856"""
    by_code = collections.defaultdict(list)
    for line in open(REG / "wikidata" / "entities.jsonl", encoding="utf-8"):
        if not line.strip():
            continue
        e = json.loads(line)
        for st in e["P429"]:
            digits = re.sub(r"\D", "", str(st["value"] or ""))
            if len(digits) == 5:
                digits = check_digit(digits)
            if st["rank"] != "deprecated" and len(digits) == 6 and check_digit(digits[:5]) == digits:
                by_code[digits].append(e)
    sites, items = {}, {}
    for code, es in by_code.items():
        alive = [e for e in es if not [s for s in e["P576"] if s["rank"] != "deprecated"]]
        es = alive or es
        urls = []
        for e in es:
            st = [s for s in e["P856"] if s["rank"] != "deprecated" and s["value"]]
            st = [s for s in st if s["rank"] == "preferred"] or st      # Wikidata の「最良の順位」と同じ
            ja = [s for s in st if not s["lang"] or J_LANG in s["lang"]]  # 日本語版か言語の指定なし
            urls += [s["value"] for s in (ja or st)]
        sites[code] = order_sites({norm_url(u) for u in urls})
        items[code] = sorted({e["id"] for e in es})
    return sites, items


def fuzzy_roma(s):
    """ローマ字の書き方の揺れ (ヘボン式と訓令式、長音) をならす"""
    s = re.sub(r"[^a-z]", "", s.lower())
    for a, b in (("mb", "nb"), ("mp", "np"), ("mm", "nm"), ("ou", "o"), ("oo", "o"), ("oh", "o"), ("uu", "u"),
                 ("shi", "si"), ("chi", "ti"), ("tsu", "tu"), ("ji", "zi"), ("sh", "s"), ("ch", "t"), ("j", "z"), ("fu", "hu")):
        s = s.replace(a, b)
    return s


def host_has_name(url, roma):
    """ホスト名に自治体のローマ字名 (先頭 4 文字) が入っているか。例: Itayanagi-machi と town.tsuruta.lg.jp は入っていない"""
    base = re.sub(r"-(shi|machi|cho|mura|son|ku)$", "", roma.lower().split()[-1]) if roma else ""
    return bool(base) and fuzzy_roma(base)[:4] in fuzzy_roma(host_of(url))


NAME_FIX = str.maketrans({"ヶ": "ケ", "竃": "竈", "桧": "檜", "舘": "館", "﨑": "崎", "髙": "高", "　": "", " ": "", "\xa0": ""})


def load_jlis(munis):
    """J-LIS の都道府県ページ: 表 (th の中の a) と地図 (area) の名前と URL"""
    cand = collections.defaultdict(list)
    for m in munis:
        cand[(m["pref"], m["name"].translate(NAME_FIX))].append(m["lg_code"])
    # 同じ県に同じ名前 (北海道の泊村: 古宇郡と北方領土の国後郡) があれば北方領土でない方
    by_name = {k: ([c for c in v if c not in NT] or v)[0] for k, v in cand.items()}
    found = {}
    for f in sorted(glob.glob(str(REG / "jlis" / "*_cms_*.html"))):
        t = open(f, encoding="utf-8", errors="replace").read()
        title = re.search(r"<title>J-LIS (.*?)</title>", t).group(1)
        pref = "北海道" if title.startswith("道") else title
        links = []
        for th in re.findall(r"<th>(.*?)</th>", t, re.S):
            for href, name in re.findall(r'<a [^>]*href="([^"]+)"[^>]*>(.*?)</a>', th, re.S):
                links.append((html.unescape(re.sub(r"<[^>]+>", "", name)), href))
        for tag in re.findall(r"<area [^>]*>", t):
            ti, hr = re.search(r'title="([^"]*)"', tag), re.search(r'href="([^"]*)"', tag)
            if ti and hr and ti.group(1) and hr.group(1):
                links.append((html.unescape(ti.group(1)), hr.group(1)))
        for name, href in links:
            code = by_name.get((pref, name.translate(NAME_FIX)))
            if code and href.startswith("http") and code not in found:
                found[code] = norm_url(href)
    return found


# ---------- 環境省 一般廃棄物処理実態調査 R6 ----------
def load_env():
    warnings.filterwarnings("ignore")
    wb = openpyxl.load_workbook(REG / "env" / "r6_shori_total_01.xlsx", read_only=True)
    ws = wb["ごみ処理概要"]
    num = lambda v: v if isinstance(v, (int, float)) else 0
    env = {}
    for r in ws.iter_rows(min_row=7, values_only=True):
        if r[1] and re.fullmatch(r"\d{5}", str(r[1])):
            env[str(r[1])] = {"name": r[2], "population": int(num(r[3])), "planned_collection_population": int(num(r[4])),
                              "self_disposal_population": int(num(r[5])), "planned_collection_t": num(r[7]),
                              "direct_haul_t": num(r[8])}
    return env


# ---------- ABR の配布ファイルの有無 ----------
def load_abr_files():
    feed = json.load(open(REG / "abr" / "dcat-us-1.1.json", encoding="utf-8"))
    have = collections.defaultdict(set)
    for ds in feed["dataset"]:
        for dist in ds.get("distribution", []):
            u = dist.get("downloadURL") or dist.get("accessURL") or ""
            m = re.search(r"/(mt_town|mt_parcel|mt_parcel_pos)/city/\1_city(\d{6})\.csv\.zip$", u)
            if m:
                have[m.group(1)].add(m.group(2))
    return have


def load_abr_town_rows():
    """ABR mt_town_all の町字の行の数 (コードごと)。索引にファイルがあっても中身が 0 行のことがある (北方領土の 6 村)"""
    return collections.Counter(r["lg_code"] for r in zip_csv_rows(REG / "abr" / "mt_town_all.csv.zip"))


# ---------- 指定離島 ----------
# N03 の面積と国交省などの面積の食い違い (町字の名前で当たった島で見たもの): 2 km² 以上の島はほぼ 3% 以内、
# 小さい島は 10〜40% ずれることがある (北九州市 馬島 1.22 倍、上天草市 湯島 1.42 倍)。
# 倉敷市 松島 (0.08 km²) の N03 のかたまりは 0.05 km² で、面積だけで探すと近くの別の岩 (0.085 km²) に当たってしまう。
BIG_KM2 = 2.0
NAME_TOL = {True: math.log(1.25), False: math.log(1.6)}    # 町字の名前で当てるとき、面積の比がこの中 (True = 2 km² 以上)
AREA_SPREAD = {True: math.log(1.04), False: math.log(2)}   # 面積だけで当てるとき、これより近いかたまりが 1 つだけのこと
GROUP_MAX = 6               # つながった島として面積の和を比べる島の数の上限
M_NAME = "町字の名前と面積"
M_GROUP = "町字の名前と面積 (N03 でつながった島をまとめて)"
M_EXACT = "町字の名前 (面積の値なし)"
M_KNOWN = "既知の位置"
M_AREA = "面積 (近い大きさのかたまりがほかに無い)"


def name_forms(s):
    """名前の書き方の揺れ (ヶ・ケ、ノ・の・之 の有無) をならした形"""
    s = s.replace("ヶ", "ケ").replace("ヵ", "ケ")
    return {s, re.sub("[ノの之]", "", s)}


def island_keys(island):
    """島の名前 (括弧の中は除く) と、末尾の「島」を除いた形。2 文字未満は使わない"""
    core = re.sub(r"\(.*?\)", "", island)
    full = {f for f in name_forms(core) if len(f) >= 2}
    base = {f for f in name_forms(core[:-1]) if len(f) >= 2} if core.endswith("島") else set()
    return full, base - full


def match_islands(desig, comps, muni_codes, towns_by_code, share):
    """指定離島を N03 の陸のかたまり (主な 5 島以外) に当てる。確かめられないものは空のまま (まちがって当てるよりよい)

      1. 町字の名前: 島の名前 (末尾の「島」を除いた形も) を含む、その自治体の町字の代表点があるかたまり。
         島の面積とかたまりの面積 (全体か、その自治体の分) の比が、2 km² 以上の島は 1/1.25〜1.25、
         それより小さい島は 1/1.6〜1.6 のときだけ当てる。
         N03 でつながっている島 (日ノ島・有福島・漁生浦島など) は面積の和で比べ、まとめて同じかたまりに当てる。
         面積の値が無い島 (父島・母島) は、名前がそのままの町字があるときだけ
      2. 既知の位置: 硫黄島・南鳥島 (no_service_areas.json と同じ位置)
      3. 面積だけ: その自治体のまだ当たっていないかたまりのうち、面積の比が食い違いの幅
         (2 km² 以上の島は 1.04 倍、それより小さい島は 2 倍) に入るものが 1 つだけで、その比が 1 の条件の中にあり、
         ほかのまだ当たっていない指定離島もそのかたまりに当たり得ないとき。
         島の名前そのものの町字があるのに 1 で当たらなかった島 (代表点が本土側・面積が合わない) はここでも当てない
    """
    off = collections.defaultdict(set)      # N03 コード → 主な島以外のかたまり
    for c, r in comps.items():
        if not r["main"]:
            for code in r["n03_codes"].split():
                off[code].add(c)

    def err(c, lg, a):
        """面積の比の対数の差。かたまり全体とその自治体の分の近い方"""
        whole = comps[c]["area_km2"]
        part = sum(share.get((c, code), 0.0) for code in muni_codes[lg])
        return min((abs(math.log(x / a)) for x in (whole, part) if x > 0), default=math.inf)

    rows = [d for d in desig if d["island"]]
    for d in desig:
        d["n03_comp_id"] = d["n03_match"] = d["n03_area_km2"] = None
    taken = set()                            # (自治体, かたまり)

    def put(d, c, how):
        d["n03_comp_id"], d["n03_match"] = c, how
        taken.add((d["lg_code"], c))

    # 1. 町字の名前
    claims, exact = collections.defaultdict(set), collections.defaultdict(set)
    named = set()                            # 島の名前そのものを含む町字がある島 (主な島の上も含む)
    for i, d in enumerate(rows):
        full, base = island_keys(d["island"])
        for code in muni_codes[d["lg_code"]]:
            for name, c in towns_by_code.get(code, []):
                forms = name_forms(name)
                if any(k in t for k in full for t in forms):
                    named.add(i)
                if c not in off[code]:
                    continue                 # 主な島の上の町字 (湖の島・本土側の代表点・同じ名前の別の所)
                if any(k in t for k in full | base for t in forms):
                    claims[(d["lg_code"], c)].add(i)
                if full & {re.sub("^大?字", "", t) for t in forms}:
                    exact[i].add(c)
    cand = []
    for (lg, c), idx in claims.items():
        idx = sorted(i for i in idx if rows[i]["area_km2"])
        for n in range(1, min(len(idx), GROUP_MAX) + 1):
            for grp in itertools.combinations(idx, n):
                a = sum(float(rows[i]["area_km2"]) for i in grp)
                e = err(c, lg, a)
                if e <= NAME_TOL[a >= BIG_KM2]:
                    cand.append((e, n, lg, c, grp))
    done = set()
    for e, n, lg, c, grp in sorted(cand):
        if (lg, c) not in taken and not done & set(grp):
            for i in grp:
                put(rows[i], c, M_NAME if n == 1 else M_GROUP)
            done |= set(grp)
    for i, cs in exact.items():
        d = rows[i]
        if not d["area_km2"] and d["n03_comp_id"] is None and len(cs) == 1 and (d["lg_code"], min(cs)) not in taken:
            put(d, min(cs), M_EXACT)

    # 2. 既知の位置 (かたまりの範囲に入るもののうちいちばん小さいもの)
    for d in rows:
        pos = DESIG_KNOWN_POS.get((d["lg_code"], d["island"]))
        if d["n03_comp_id"] is None and pos:
            lon, lat = pos
            inside = [c for code in muni_codes[d["lg_code"]] for c in off[code]
                      if float(comps[c]["min_lon"]) <= lon <= float(comps[c]["max_lon"])
                      and float(comps[c]["min_lat"]) <= lat <= float(comps[c]["max_lat"])]
            if inside:
                put(d, min(inside, key=lambda c: comps[c]["area_km2"]), M_KNOWN)

    # 3. 面積だけ (同じ自治体の中で同時に決める。決めた結果で次を決めることはしない)
    left = collections.defaultdict(list)
    for i, d in enumerate(rows):
        if d["n03_comp_id"] is None and d["area_km2"]:
            left[d["lg_code"]].append((i, d, float(d["area_km2"])))
    for lg, ds in left.items():
        cs = {c for code in muni_codes[lg] for c in off[code] if (lg, c) not in taken}
        picks = []
        for i, d, a in ds:
            if i in named:
                continue
            near = [c for c in cs if err(c, lg, a) <= AREA_SPREAD[a >= BIG_KM2]]
            if len(near) != 1 or err(near[0], lg, a) > NAME_TOL[a >= BIG_KM2]:
                continue                     # 当たり得るかたまりが無いか、2 つ以上あるか、面積が離れすぎ
            if any(err(near[0], lg, ao) <= AREA_SPREAD[ao >= BIG_KM2] for j, _o, ao in ds if j != i):
                continue                     # ほかの指定離島もこのかたまりに当たり得る
            picks.append((d, near[0]))
        for d, c in picks:
            put(d, c, M_AREA)
    for d in desig:
        d["n03_muni_area_km2"] = None
        if d["n03_comp_id"]:
            c = d["n03_comp_id"]
            d["n03_area_km2"] = comps[c]["area_km2"]
            d["n03_muni_area_km2"] = round(sum(share.get((c, code), 0.0) for code in muni_codes[d["lg_code"]]), 4)


# ---------- 答えない所 ----------
def bbox_of(rows):
    return [min(float(r["min_lon"]) for r in rows), min(float(r["min_lat"]) for r in rows),
            max(float(r["max_lon"]) for r in rows), max(float(r["max_lat"]) for r in rows)]


def comp_entry(c, name=None):
    e = {"comp_id": int(c["comp_id"]), "area_km2": float(c["area_km2"]),
         "bbox": [float(c[k]) for k in ("min_lon", "min_lat", "max_lon", "max_lat")],
         "rep_point": [float(c["rep_lon"]), float(c["rep_lat"])]}
    if name:
        e["name"], e["name_source"] = name
    if c.get("towns"):
        e["abr_towns"] = c["towns"].split("/")[:10]
    return e


def near(c, lon, lat, deg):
    return abs(float(c["rep_lon"]) - lon) <= deg and abs(float(c["rep_lat"]) - lat) <= deg


def island_names(cs, refs, deg):
    """かたまりに島の名前を付ける: ABR 町字の名前 (「…島」) → 無ければ既知の位置にいちばん近い 0.01km² 以上のかたまり"""
    names = {}
    for c in cs:
        isl = [t for t in c["towns"].split("/") if t.endswith("島")] if c["towns"] else []
        if isl:
            names[c["comp_id"]] = ("・".join(dict.fromkeys(isl)), "ABR 町字")
    for nm, lon, lat in refs or []:
        if nm in {v[0] for v in names.values()}:
            continue
        cand = [c for c in cs if near(c, lon, lat, deg) and c["comp_id"] not in names and c["area_km2"] >= 0.01]
        if cand:
            c = min(cand, key=lambda c: (float(c["rep_lon"]) - lon) ** 2 + (float(c["rep_lat"]) - lat) ** 2)
            names[c["comp_id"]] = (nm, "位置から推定")
    return names


# 既知の島の位置 (名前を付けるためだけに使う。ABR 町字の名前がある島はそちらを優先)
TAKESHIMA = [("男島 (西島)", 131.8644, 37.2422), ("女島 (東島)", 131.8697, 37.2417)]
SENKAKU = [("魚釣島", 123.4766, 25.7432), ("北小島", 123.5416, 25.7302), ("南小島", 123.5556, 25.7231),
           ("久場島", 123.6828, 25.9259), ("大正島", 124.5594, 25.9225)]
HABOMAI = [("貝殻島", 145.883, 43.383)]    # ほかの島 (志発島・水晶島・多楽島・勇留島・秋勇留島) は ABR 町字の名前で付く
IWOTO = [("硫黄島", 141.3195, 24.7793)]
MINAMITORI = [("南鳥島", 153.9806, 24.2892)]
DESIG_KNOWN_POS = {("134210", "硫黄島"): IWOTO[0][1:], ("134210", "南鳥島"): MINAMITORI[0][1:]}   # 指定離島を当てるのにも使う
MIN_LIST_KM2 = 0.001     # これより小さく町字も無い岩は components に 1 つずつは書かない (グループの bbox には入る)


def build_no_service(munis_by_code, comps, parts, profile, main_comps):
    areas = []
    by_code = collections.defaultdict(list)
    for p in parts:
        by_code[p["n03_code"]].append(p)

    def groups(code):
        """そのコードのポリゴンを陸のかたまりごとにまとめる"""
        g = collections.defaultdict(list)
        for p in by_code[code]:
            g[p["comp"]].append(p)
        out = []
        for c, ps in sorted(g.items(), key=lambda kv: -sum(float(p["area_km2"]) for p in kv[1])):
            big = max(ps, key=lambda p: float(p["area_km2"]))
            out.append({"comp_id": int(c), "on_main_island": main_comps.get(c),
                        "area_km2": round(sum(float(p["area_km2"]) for p in ps), 6), "bbox": bbox_of(ps),
                        "rep_point": [float(big["rep_lon"]), float(big["rep_lat"])]})
        return out

    for code in sorted(NT):
        m = munis_by_code[code]
        pr = profile[code[:5]]
        areas.append({
            "id": f"northern_territory_{code}", "type": "northern_territory", "lg_code": code, "n03_code": code[:5],
            "pref": m["pref"], "name": m["name"], "gps2city_lg_code": code,
            "reason": "北方領土の村。ロシアが実効支配していて日本の行政が及ばず、ごみ収集の情報も無い",
            "area_km2": float(pr["area_km2"]), "bbox": bbox_of([pr]),
            "rep_point": [float(pr["rep_lon"]), float(pr["rep_lat"])], "pieces": groups(code[:5])})

    mitei = sorted(c for c in by_code if c.endswith("000"))
    assert set(mitei) == SHOZOKU_MITEI, mitei
    for code in mitei:
        pr = profile[code]
        areas.append({
            "id": f"shozoku_mitei_{code}", "type": "unassigned_land", "lg_code": None, "n03_code": code,
            "pref": pr["name"].removesuffix("所属未定地"), "name": pr["name"],
            "gps2city_lg_code": check_digit(code),
            "reason": "所属未定地 (どの市区町村にも属さない土地)。gps2city は都道府県のコード "
                      f"{check_digit(code)} を返すが、市区町村の収集の決まりは無い",
            "area_km2": float(pr["area_km2"]), "bbox": bbox_of([pr]),
            "rep_point": [float(pr["rep_lon"]), float(pr["rep_lat"])], "pieces": groups(code)})

    def island_group(gid, kind, lg, name, reason, sel, refs=None, deg=0.03):
        cs = [c for c in comps.values() if lg[:5] in c["n03_codes"].split() and sel(c)]
        assert cs, gid
        names = island_names(cs, refs, deg)
        listed = [c for c in cs if c["area_km2"] >= MIN_LIST_KM2 or c["comp_id"] in names or int(c["n_towns"]) > 0]
        ents = [comp_entry(c, names.get(c["comp_id"])) for c in sorted(listed, key=lambda c: -float(c["area_km2"]))]
        m = munis_by_code[lg]
        areas.append({"id": gid, "type": kind, "lg_code": lg, "n03_code": lg[:5], "pref": m["pref"],
                      "municipality": m["name"], "name": name, "gps2city_lg_code": lg, "reason": reason,
                      "area_km2": round(sum(c["area_km2"] for c in cs), 4), "bbox": bbox_of(cs),
                      "rep_point": ents[0]["rep_point"], "n_components": len(cs),
                      "small_components_not_listed": len(cs) - len(ents), "components": ents})

    island_group("takeshima", "unadministered_island", "325287", "竹島",
                 "竹島。韓国が実効支配していて日本の行政が及ばない (gps2city は隠岐の島町を返す)",
                 lambda c: float(c["rep_lon"]) < 132.5, TAKESHIMA, 0.01)     # 隠岐 (東経 133° 付近) より西
    island_group("senkaku", "uninhabited_island", "472077", "尖閣諸島",
                 "尖閣諸島。住民がいない (gps2city は石垣市を返す)",
                 lambda c: float(c["rep_lat"]) > 25.4, SENKAKU, 0.02)
    island_group("iwoto", "no_resident_island", "134210", "硫黄島",
                 "硫黄島。住民はいない (自衛隊などの職員だけ)。村のごみ収集は無い (gps2city は小笠原村を返す)",
                 lambda c: near(c, 141.3195, 24.7793, 0.1), IWOTO, 0.05)
    island_group("minamitorishima", "no_resident_island", "134210", "南鳥島",
                 "南鳥島。住民はいない (気象庁・自衛隊などの職員だけ)。村のごみ収集は無い (gps2city は小笠原村を返す)",
                 lambda c: near(c, 153.9806, 24.2892, 0.05), MINAMITORI, 0.05)
    island_group("okinotorishima", "uninhabited_island", "134210", "沖ノ鳥島",
                 "沖ノ鳥島。人の住めない岩礁 (gps2city は小笠原村を返す)",
                 lambda c: near(c, 136.0811, 20.4256, 0.1))
    island_group("habomai", "unadministered_island", "012238", "歯舞群島",
                 "歯舞群島。根室市に属するが、ロシアが実効支配していて日本の行政が及ばない (gps2city は根室市を返す)",
                 lambda c: float(c["min_lon"]) >= 145.84, HABOMAI, 0.02)     # 納沙布岬 (東経 145.82°) より東
    return areas


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    jlis_fallback = "--no-jlis-fallback" not in sys.argv[1:]
    munis, wards = load_soumu()
    by_code = {m["lg_code"]: m for m in munis}

    # ---- 件数の確かめ (調査と同じ入力なら同じ数になるはず)
    kinds = collections.Counter(m["kind"] for m in munis)
    print("自治体コード", len(munis), dict(kinds), "政令市", len(wards), "区", sum(map(len, wards.values())))
    assert len(munis) == 1747 and kinds["北方領土"] == 6
    assert (kinds["市"], kinds["町"], kinds["村"], kinds["特別区"]) == (792, 743, 183, 23)
    assert len(wards) == 20 and sum(map(len, wards.values())) == 171
    abr_city = list(zip_csv_rows(REG / "abr" / "mt_city_all.csv.zip"))
    abr_codes = {r["lg_code"] for r in abr_city if r["lg_code"][2:] != "0000"}
    ward_codes = {w for ws in wards.values() for w in ws}
    print("ABR mt_city との差:", sorted(abr_codes ^ ({m["lg_code"] for m in munis} | ward_codes)) or "なし")
    county = {r["lg_code"]: r["county"] for r in abr_city}
    roma = {r["lg_code"]: r["city_roma"] for r in abr_city if not r["ward"]}

    # ---- 島
    cls = json.load(open(WORK / "muni_island_class.json", encoding="utf-8"))
    main_comps = {int(k): v for k, v in cls["main_components"].items()}
    profile = {r["n03_code"]: r for r in read_csv(OUT / "municipality_islands.csv")}
    comps = {int(r["comp_id"]): dict(r, area_km2=float(r["area_km2"])) for r in read_csv(WORK / "n03_components.csv")}
    muni_comps = collections.defaultdict(list)       # 自治体 → 町字のある島 (主な島以外)
    muni_codes = collections.defaultdict(set)        # 自治体 → N03 コード (政令市は区のコード)
    for r in profile.values():
        muni_codes[r["muni_lg_code"]].add(r["n03_code"])
        for c in r["addressed_islands"].split():
            muni_comps[r["muni_lg_code"]].append(int(c))
    parts = read_csv(WORK / "n03_parts.csv")
    share = collections.Counter()                    # (かたまり, N03 コード) → 面積
    for p in parts:
        share[(int(p["comp"]), p["n03_code"])] += float(p["area_km2"])
    desig = read_csv(HERE / "designated_islands.csv")
    want = {code for d in desig for code in muni_codes[d["lg_code"]]}
    towns_by_code = collections.defaultdict(list)    # N03 コード → 町字 (名前, 代表点のあるかたまり)
    for t in csv.DictReader(open(WORK / "town_components.csv", encoding="utf-8")):
        if t["name"] and t["lg_code"][:5] in want:
            towns_by_code[t["lg_code"][:5]].append((t["name"], int(t["comp"])))
    match_islands(desig, comps, muni_codes, towns_by_code, share)
    dz = collections.defaultdict(list)
    for d in desig:
        assert d["lg_code"] in by_code, d
        dz[d["lg_code"]].append(d)
    law_count = collections.Counter((d["law"], d["inhabited"]) for d in desig)
    print("指定離島の行", len(desig), dict(law_count), "指定のある自治体", len(dz))

    # ---- 公式サイト・環境省・ABR
    wd_sites, wd_items = load_wikidata()
    jlis = load_jlis(munis)
    env = load_env()
    abr_files = load_abr_files()
    town_rows = load_abr_town_rows()
    no_service = build_no_service(by_code, comps, parts, profile, main_comps)
    ns_by_lg = collections.defaultdict(list)
    for a in no_service:
        if a["lg_code"] and a["type"] != "northern_territory":
            ns_by_lg[a["lg_code"]].append(a["name"])

    rows, stats, jlis_diff = [], collections.Counter(), []
    for m in munis:
        code, notes = m["lg_code"], []
        units = wards.get(code) or [code]
        isl_class = cls["classes"][code]
        ds = dz.get(code, [])
        laws = sorted({d["law"] for d in ds})
        inhabited = [{"island": d["island"], "law": d["law"], "inhabited": d["inhabited"],
                      "population_2020": int(d["population_2020"]) if d["population_2020"] else None,
                      "n03_comp_id": d["n03_comp_id"], "n03_match": d["n03_match"], "note": d["note"] or None}
                     for d in ds if d["inhabited"] in ("yes", "staff_only")]
        # ABR: 町字は索引にファイルがあり、mt_town_all に行もあること。地番と地番の位置は索引にファイルがあるかだけ
        has = {"mt_town": lambda u: u in abr_files["mt_town"] and town_rows[u] > 0,
               "mt_parcel": lambda u: u in abr_files["mt_parcel"],
               "mt_parcel_pos": lambda u: u in abr_files["mt_parcel_pos"]}
        cov = {k: all(f(u) for u in units) for k, f in has.items()}
        cov["mt_town_rows"] = sum(town_rows[u] for u in units)
        missing = {k: [u for u in units if not f(u)] for k, f in has.items()}
        missing = {k: v for k, v in missing.items() if v and len(v) < len(units)}
        if missing:
            cov["missing_units"] = missing
        town_file_empty = any(u in abr_files["mt_town"] and not town_rows[u] for u in units)
        # 環境省
        e = env.get(code[:5])
        env_out = None
        if e:
            flags = []
            if e["self_disposal_population"] > 0:
                flags.append("no_public_collection_partial")
            if e["planned_collection_t"] == 0:
                if m["kind"] == "特別区":
                    notes.append("環境省の調査では 23 区の計画収集量は「東京都23区」の合計行 (13100) にだけあり、区ごとは 0")
                else:
                    flags.append("collection_by_drop_off_only")
            env_out = {k: v for k, v in e.items() if k != "name"}
            env_out["flags"] = flags
            stats.update(flags)
        elif m["kind"] != "北方領土":
            stats["env_missing"] += 1
        # 公式サイト: Wikidata。J-LIS は突き合わせに使い、URL は Wikidata に自治体のサイトが無いときだけ使う
        wd = wd_sites.get(code, [])
        j = jlis.get(code)
        if m["kind"] == "北方領土":
            # 村の役場は無い。Wikidata に URL があっても別の村 (古宇郡泊村など) のものなので使わない
            wd, j = [], None
        # 1. ホスト名に自治体名が無く J-LIS とも違う URL は荒らしや取り違えとみて外す
        #    (例: 御浜町の tiktocoins.info、板柳町に入っていた鶴田町の URL)
        dropped = [u for u in wd if not host_has_name(u, roma.get(code, "")) and not (j and host_of(u) == host_of(j))]
        sites, src = [u for u in wd if u not in dropped], "wikidata"
        # 2. 自治体のホストらしい URL (lg.jp・city./town./vill.) が 1 つも無く、J-LIS が lg.jp を載せていればそれを使う
        #    (lg.jp は自治体しか登録できない)。残らないとき (板柳町) も J-LIS
        jlis_used = bool(j) and not any(gov_like(u) for u in sites) and (not sites or host_of(j).endswith(".lg.jp"))
        if jlis_used:
            sites = order_sites(sites + [j])
        # 3. 自治体のホストらしい URL があれば、それ以外 (.org・.com など。観光協会のことが多い) は外す
        #    (例: 白川村の shirakawa-go.org、南会津町の minamiaizu.org)
        non_muni = [u for u in sites if not gov_like(u)] if any(gov_like(u) for u in sites) else []
        sites = [u for u in sites if u not in non_muni]
        if jlis_used:
            src = "jlis"                     # Wikidata の URL は自治体のホストらしくないので 3 で外れる
        if not sites and dropped:
            sites, dropped = dropped, []
        if not sites:
            src = None
            stats["no_site_" + m["kind"]] += 1
        same = (host_of(j) in {host_of(u) for u in wd}) if (j and wd) else None
        # 4. Wikidata と J-LIS のホストが違えば (Wikidata の URL が古い・移転前のことがある。北本市など)、J-LIS の URL を
        #    2 つ目の公式サイトとして足す (印 jlis_fallback)。探すとき (discover.py) は、1 つ目につながらないか収集日の資料が
        #    見つからなければ 2 つ目を見る。自治体のホストらしい URL どうしのときだけ (観光協会らしいものと混ぜない)
        fallback = bool(jlis_fallback and same is False and not jlis_used and sites and gov_like(j)
                        and all(gov_like(u) for u in sites) and host_of(j) not in {host_of(u) for u in sites})
        if fallback:
            sites = sites + [j]
        # J-LIS の URL そのものは、使ったときと 2 つ目に足したときだけ official_sites に入る (ほかは同じホストかどうかだけ残す)
        check = {"wikidata_items": wd_items.get(code, []), "wikidata_dropped": dropped, "non_municipal_dropped": non_muni,
                 "jlis_used": jlis_used, "jlis_same_host": same,
                 "needs_review": bool(dropped) or bool(non_muni) or jlis_used or same is False}
        if fallback:
            check["jlis_fallback"] = True    # official_sites の最後の URL は J-LIS のもの
        stats["jlis_fallback"] += fallback
        stats["wikidata_dropped"] += len(dropped)
        stats["non_municipal_dropped"] += len(non_muni)
        stats["needs_review"] += check["needs_review"]
        stats["jlis_" + {True: "same", False: "differs", None: "n/a"}[check["jlis_same_host"]]] += 1
        if same is False:
            jlis_diff.append(f"{code} {m['name']}: Wikidata {' '.join(wd)} / J-LIS {j}")
        # 注記
        if m["kind"] == "北方領土":
            notes.append("北方領土の村。gomisute-gps は答えない (no_service_areas.json)")
        if town_file_empty:
            notes.append("ABR の配布索引には町字のファイルがあるが、mt_town_all に町字の行が無い")
        if jlis_used:
            notes.append("公式サイトは J-LIS の一覧から取った (Wikidata に自治体のサイトが無い)。人が確かめること")
        if fallback:
            notes.append("公式サイトの 2 つ目は J-LIS の URL (Wikidata とホストが違う。1 つ目が古いときのため)。人が確かめること")
        for nm in ns_by_lg.get(code, []):
            notes.append(f"{nm} は答えない (no_service_areas.json)")
        if laws and isl_class in ("main_only", "main+unaddressed_islets"):
            names = "・".join(d["island"] for d in ds if d["island"] and d["n03_comp_id"] is None) or "指定地域"
            notes.append(f"指定離島 ({names}) があるが、N03 と ABR 町字では「町字のある島」にならない"
                         "(湖の島、町字の代表点が本土側、または無人島だけ)")
        if isl_class == "all_island" and not laws and m["kind"] != "北方領土":
            notes.append("全域が主な 5 島の外にあるが、離島の法律の指定は無い (橋でつながる島など)")
        rows.append({
            "lg_code": code, "pref": m["pref"], "county": county.get(code, ""), "name": m["name"], "kind": m["kind"],
            "wards": wards.get(code, []), "island_class": isl_class, "island_law": laws,
            "inhabited_islands": inhabited,
            "n03_addressed_islands": sorted(muni_comps.get(code, [])),
            "abr_parcel_coverage": cov, "env_r6": env_out,
            "official_sites": sites, "official_sites_source": src, "official_sites_check": check,
            "coverage": {"answer_type": "dates" if code in DATES else "none_yet",
                         "source": "gomisute-gps data" if code in DATES else None},
            "no_service": code in NT, "notes": notes,
        })

    # ---- 書き出し
    OUT.mkdir(parents=True, exist_ok=True)
    manifest = json.load(open(REG / "manifest.json", encoding="utf-8"))
    eff = [r for r in rows if r["kind"] != "北方領土"]
    summary = {
        "municipality_codes": len(rows), "in_effect": len(eff),
        "kinds": dict(collections.Counter(r["kind"] for r in rows)),
        "designated_cities": len(wards), "wards": sum(map(len, wards.values())),
        "island_class_in_effect": dict(collections.Counter(r["island_class"] for r in eff)),
        "island_law_municipalities": len([r for r in rows if r["island_law"]]),
        "with_designated_inhabited_island": len([r for r in rows if r["inhabited_islands"]]),
        "island_law_by_class": dict(collections.Counter(r["island_class"] for r in rows if r["island_law"])),
        "designated_islands": {f"{k[0]}:{k[1] or 'none'}": v for k, v in sorted(law_count.items())},
        "designated_islands_matched_to_n03": dict(collections.Counter(d["n03_match"] or "none" for d in desig if d["island"])),
        "env_flags": {k: v for k, v in stats.items() if k in ("no_public_collection_partial", "collection_by_drop_off_only", "env_missing")},
        "official_sites": {"wikidata": sum(r["official_sites_source"] == "wikidata" for r in rows),
                           "jlis": sum(r["official_sites_source"] == "jlis" for r in rows),
                           "none_in_effect": [r["lg_code"] + " " + r["name"] for r in eff if not r["official_sites"]],
                           "none_northern_territories": sum(not r["official_sites"] for r in rows if r["kind"] == "北方領土"),
                           "urls": sum(len(r["official_sites"]) for r in rows),
                           "lg_jp_first": sum(bool(r["official_sites"]) and host_of(r["official_sites"][0]).endswith(".lg.jp") for r in rows),
                           "jlis_same_host": stats["jlis_same"], "jlis_differs": stats["jlis_differs"],
                           "jlis_not_compared": stats["jlis_n/a"], "jlis_fallback_added": stats["jlis_fallback"],
                           "wikidata_urls_dropped": stats["wikidata_dropped"],
                           "non_municipal_urls_dropped": stats["non_municipal_dropped"], "needs_review": stats["needs_review"]},
        "abr_all_units": {k: sum(r["abr_parcel_coverage"][k] for r in eff) for k in ("mt_town", "mt_parcel", "mt_parcel_pos")},
    }
    meta = {"generated_at": datetime.date.today().isoformat(), "generated_by": "nationwide/build_registry.py",
            "n03": cls["n03"], "summary": summary,
            "sources": {k: {"url": v["url"], "sha256": v["sha256"], "last_modified": v.get("last_modified")}
                        for k, v in sorted(manifest.items()) if not k.startswith(("abr/mt_town_pos/", "jlis/"))}}
    json.dump({"meta": meta, "municipalities": rows}, open(OUT / "municipalities.json", "w", encoding="utf-8"),
              ensure_ascii=False, indent=1)
    flat = ["lg_code", "pref", "county", "name", "kind", "wards", "island_class", "island_law", "inhabited_islands",
            "n03_addressed_islands", "abr_town", "abr_parcel", "abr_parcel_pos", "env_population",
            "env_self_disposal_population", "env_planned_collection_t", "env_flags", "official_sites",
            "official_sites_source", "jlis_same_host", "answer_type", "coverage_source", "no_service", "notes"]
    with open(OUT / "municipalities.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(flat)
        for r in rows:
            e = r["env_r6"] or {}
            w.writerow([r["lg_code"], r["pref"], r["county"], r["name"], r["kind"], " ".join(r["wards"]),
                        r["island_class"], " ".join(r["island_law"]),
                        " ".join(i["island"] + ("(職員のみ)" if i["inhabited"] == "staff_only" else "") for i in r["inhabited_islands"]),
                        " ".join(map(str, r["n03_addressed_islands"])),
                        int(r["abr_parcel_coverage"]["mt_town"]), int(r["abr_parcel_coverage"]["mt_parcel"]),
                        int(r["abr_parcel_coverage"]["mt_parcel_pos"]), e.get("population", ""),
                        e.get("self_disposal_population", ""), e.get("planned_collection_t", ""),
                        " ".join(e.get("flags", [])), " ".join(r["official_sites"]), r["official_sites_source"] or "",
                        "" if r["official_sites_check"]["jlis_same_host"] is None else int(r["official_sites_check"]["jlis_same_host"]),
                        r["coverage"]["answer_type"], r["coverage"]["source"] or "", int(r["no_service"]),
                        " / ".join(r["notes"])])
    json.dump({"meta": {"generated_at": meta["generated_at"], "generated_by": meta["generated_by"], "n03": cls["n03"],
                        "how_to_use": "northern_territory と unassigned_land は、gps2city の lg_code が gps2city_lg_code と同じなら"
                                      "場所によらず答えない (所属未定地では県のコードが返る)。島のグループは、gps2city の lg_code が"
                                      "グループの lg_code と同じで、点がグループの bbox に入れば答えない (bbox に住民のいる土地は入らない)。"
                                      "components は島ごとの範囲 (0.001km² 未満の岩は省略)"},
               "areas": no_service}, open(OUT / "no_service_areas.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    # 指定離島と N03 のかたまりの対応 (確かめやすいように全部の行)
    cols = ["island", "pref", "municipality", "lg_code", "law", "inhabited", "area_km2", "n03_comp_id", "n03_area_km2",
            "n03_muni_area_km2", "n03_match"]
    with open(OUT / "designated_islands_n03.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(cols)
        for d in desig:
            w.writerow(["" if d.get(k) is None else d[k] for k in cols])
    print("J-LIS とホストが違う (人が確かめる。自治体のホストらしい URL どうしなら J-LIS の URL を 2 つ目に足す):", *jlis_diff, sep="\n  ")
    print(json.dumps(summary, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
