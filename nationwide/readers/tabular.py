"""型 A: CSV・Excel の読み取り (DESIGN.md §3.5 A)

  python nationwide/readers/tabular.py <CSV・Excel のファイルか URL (キャッシュから)>    # 見出しの行・列の役目・行の読み

- 文字コードは decode.py の 1 つの判定 (BOM → HTTP の charset → UTF-8 → cp932。読めなければ止める)。見出しと BOM のぶつかりは記録に残す
- 見出しの行は中身で探す: 「町名」「地区」「曜日」「品目」・品目の名前などの見出しの言葉を 2 つ以上含む最初の行
  (世田谷の CSV は列の見出しの上に題の行がある)
- Excel は結合したマスを広げ (左上の値で埋める)、隠しシートを外す。シートの名前に期間と合わない年の印があるもの
  (狭山の 2025 年の「柏原」) は外す
- 行の読み (地区・品目・規則) は HTML の表と同じ (html_tables.read_grid)。2 つ目の読み (LLM) には、先頭と無作為の行の見本を渡す
"""
import csv, io, random, re, sys

import categories, decode, html_tables, version

HEADER_WORDS = re.compile(r"町名|地区|地域|区域|丁目|曜日|品目|種別|収集日|50音|五十音")


def read_csv(body, content_type=""):
    """CSV の bytes → (行のリスト, 文字コードの記録)"""
    d = decode.decode(body, content_type, html=False)
    text = d.text
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    return rows, d.record()


def _xlsx_sheets(body):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [[("" if v is None else str(v)).strip() for v in r] for r in ws.iter_rows(values_only=True)]
        for rng in ws.merged_cells.ranges:                  # 結合したマスは左上の値で埋める
            v = rows[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(rows) else ""
            for r in range(rng.min_row - 1, min(rng.max_row, len(rows))):
                for c in range(rng.min_col - 1, min(rng.max_col, len(rows[r]))):
                    rows[r][c] = v
        out.append({"sheet": ws.title, "hidden": ws.sheet_state != "visible", "rows": rows})
    return out


def _xls_sheets(body):
    import xlrd
    # 結合したマス (merged_cells) は formatting_info=True のときだけ読まれる (xlrd は False だと XL_MERGEDCELLS を飛ばす)
    wb = xlrd.open_workbook(file_contents=body, formatting_info=True)
    out = []
    for sh in wb.sheets():
        rows = [[str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)] for r in range(sh.nrows)]
        for (r0, r1, c0, c1) in getattr(sh, "merged_cells", []) or []:
            v = rows[r0][c0]
            for r in range(r0, r1):
                for c in range(c0, c1):
                    rows[r][c] = v
        out.append({"sheet": sh.name, "hidden": bool(getattr(sh, "visibility", 0)), "rows": rows})
    return out


def read_excel(body, period=None):
    """Excel の bytes → 読むシートのリスト [{"sheet", "rows"}] と外したシートの記録"""
    sheets = _xlsx_sheets(body) if body[:2] == b"PK" else _xls_sheets(body)
    keep, skipped = [], []
    for s in sheets:
        if s["hidden"]:
            skipped.append({"sheet": s["sheet"], "why": "hidden"})
            continue
        keep.append(s)
    keep, sk = drop_old_sheets(keep, period)
    return keep, skipped + sk


def drop_old_sheets(sheets, period):
    """期間と合わない年の印がシートの名前にあるシート (狭山の 2025 年の「柏原」) を外す → (残すシート, 外した記録)。
    期間が決まっていなければ何も外さない (呼ぶ側が期間を決めてから、もう一度呼ぶ)"""
    if not period or not period.get("valid_to"):
        return list(sheets), []
    keep, skipped = [], []
    for s in sheets:
        toks = [t for t in version.year_tokens(s["sheet"]) if t["kind"] in ("fy", "span")]
        if toks and not any(version.token_matches(t, period["valid_from"], period["valid_to"]) for t in toks):
            skipped.append({"sheet": s["sheet"], "why": "year_token_outside_period"})
            continue
        keep.append(s)
    return keep, skipped


def header_index(rows):
    """見出しの行の番号 (見出しの言葉か品目の名前を 2 つ以上含む最初の行)。無ければ None"""
    for i, r in enumerate(rows[:30]):
        hits = sum(1 for c in r if c and (HEADER_WORDS.search(c) or categories.is_category_header(c)))
        if hits >= 2:
            return i
    return None


# 題の行の年度 (「令和8年度」「2026年度」「令和8年度（2026年度）」) と、題に使う名詞の閉じた並び
_TITLE_YEAR = re.compile(r"(?:令和|平成)\s*(?:\d+|元)\s*年度?|\d{4}\s*年度?|年度")
_TITLE_NOUNS = re.compile(r"家庭|ごみ|ゴミ|資源物?|分別|収集|日程|カレンダー|スケジュール|一覧|表|日|の|[\s・()（）【】「」\[\]]")


def title_row(text, own_names=()):
    """見出しより上の行が題の行か (閉じた形): 自治体の名前と年度を除いた残りが、題の名詞 (家庭・ごみ・資源・分別・収集・日程・
    カレンダー・一覧・表・日) とかっこだけのもの (「令和8年度 ごみ収集日」「2026年度 家庭ごみ収集カレンダー」)。
    「資源は月1回」「ごみは月2回」「収集日は第1日曜日」「資源ごみ 第1・第3」は題でない (注の行として止める)。題の行は期間の手がかりにだけ使う"""
    t = html_tables.nfkc(html_tables._own_stripped(text, own_names))
    return bool(t.strip()) and not _TITLE_NOUNS.sub("", _TITLE_YEAR.sub("", t))


def read_rows(rows, lg_code=None, nth_def=None, own_names=()):
    """行のリスト → 読み (html_tables.read_grid と同じ形。行の番号は元の行の番号)。
    見出しより上の行のうち、題の行 (title_row) でないもの (「※第5週は収集しません,,」「8月は収集しません」) は表の注の行 (notes) にする:
    呼ぶ側が表のすべての系列の注として関門に通す (読めなければ止める)。CSV・Excel には本文の文の洗い出しの網が無いので、黙って落とさない"""
    h = header_index(rows)
    if h is None:
        return {"orientation": "other", "why": "見出しの行が見つからない", "header_index": None}
    rd = html_tables.read_grid(rows[h:], 1, None, nth_def, lg_code)
    sh = lambda pos: (pos[0] + h, pos[1]) if isinstance(pos, tuple) else pos
    for r in rd.get("rows", []):
        r["row"] += h                                       # 元の行の番号に戻す
        r["cells"] = {k: ([sh(x) for x in v] if isinstance(v, list) else sh(v)) for k, v in r["cells"].items()}
        oc = r.get("other_cells")
        if isinstance(oc, dict):
            r["other_cells"] = {k: [(a + h, b, c) for a, b, c in v] for k, v in oc.items()}
        elif oc:
            r["other_cells"] = [(a + h, b, c) for a, b, c in oc]
        for k in ("src_rows",):
            if r.get(k):
                r[k] = [x + h for x in r[k]]
        if r.get("district_cells"):
            r["district_cells"] = [sh(x) for x in r["district_cells"]]
    for n in rd.get("notes", []):
        n["row"] += h
        n["cells"] = [sh(x) for x in n["cells"]]
    rd["mid_heads"] = [x + h for x in rd.get("mid_heads", [])]
    rd["header_index"] = h
    titles, pre = [], []
    for i, r in enumerate(rows[:h]):
        t = " ".join(c for c in r if c)
        if not t:
            continue
        titles.append(t)                                    # 年の印は、どの行のものも期間の手がかりにする (今までどおり)
        if not title_row(t, own_names):
            pre.append({"row": i, "text": "\n".join(dict.fromkeys(c for c in r if c)), "cells": [(i, j) for j, c in enumerate(r) if c],
                        "above_header": True})
    rd["title_rows"] = titles
    if rd.get("orientation") != "other":
        rd["notes"] = pre + rd.get("notes", [])
    return rd


def note_rows(reading):
    """注の行・途中の見出しの行の元の行の番号 (2 つ目の読みにも見せる)"""
    return sorted({n["row"] for n in reading.get("notes") or [] if n["row"] >= 0} | set(reading.get("mid_heads") or []))


def sample_rows(reading, n_head_row, k=20, seed=0):
    """2 つ目の読みに渡す行の見本: 先頭の k 行と、残りから無作為の k 行 (seed を決めて、同じ見本になるように)。
    品目の行・曜日の行の表は、読みの行がまとめた元の行 (src_rows) すべて"""
    rows = sorted({x for r in reading.get("rows", []) for x in (r.get("src_rows") or [r["row"]])})
    head = rows[:k]
    rest = rows[k:]
    rnd = random.Random(seed)
    pick = sorted(rnd.sample(rest, min(k, len(rest)))) if rest else []
    return head + pick


def as_table_html(rows, keep_rows, header_row):
    """CSV の行 → 2 つ目の読みに渡す表の形 (タグは td だけ。CSV には しるし が無い)。data-row は元の行の番号。
    見出しの行をいちばん上に置き、あとは keep_rows の順 (見出しより上の注の行も、ここに入れて見せる)"""
    import html
    out = ["<table>"]
    for i in [header_row] + list(keep_rows):
        cells = "".join(f"<td>{html.escape(c)}</td>" for c in rows[i])
        out.append(f'<tr data-row="{i}">{cells}</tr>')
    out.append("</table>")
    return "\n".join(out)


def main(argv=None):
    import argparse, pathlib
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="CSV・Excel の見出しの行・列の役目・行の読み")
    ap.add_argument("src")
    a = ap.parse_args(argv)
    if a.src.startswith("http"):
        import fetch_docs
        body = fetch_docs.cached(a.src)
        if body is None:
            print("キャッシュに無い:", a.src)
            return 1
    else:
        body = pathlib.Path(a.src).read_bytes()
    if body[:2] == b"PK" or body[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        sheets, skipped = read_excel(body)
        print("外したシート:", skipped)
    else:
        rows, rec = read_csv(body)
        print("文字コード:", rec)
        sheets = [{"sheet": "csv", "rows": rows}]
    for s in sheets:
        rd = read_rows(s["rows"])
        print(f"シート {s['sheet']} 見出しの行={rd.get('header_index')} 向き={rd['orientation']} 役目={rd.get('roles')}")
        for c, v in rd.get("categories", {}).items():
            print("   品目", repr(c), v["coarse"], v["mode"])
        for r in rd.get("rows", [])[:5]:
            print("   行", r["row"], repr(r["district_text"]), {k: (p.status, p.irs) for k, p in r["rules"].items()}, r.get("calendar_links"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
