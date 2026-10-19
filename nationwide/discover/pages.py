"""取ったページ・資料を読む (HTML のリンク・題名・見出し・表、PDF の文字の有無、CSV・Excel・iCalendar の中身の始め)

外部の部品は使えれば使う: PDF は PyMuPDF (fitz)、無ければ pdfplumber。Excel は openpyxl。どれも無ければ「見ていない」と記録する。
PDF・Excel はキャッシュのファイルを bytes で読んでから開く (誤りの文に手元の場所が入らないように)。
"""
import csv, html, html.parser, io, pathlib, re, unicodedata, urllib.parse

from fetcher import scrub

try:
    import fitz                       # PyMuPDF
    try:
        fitz.TOOLS.mupdf_display_errors(False)      # 壊れた PDF の警告で画面を埋めない
    except Exception:
        pass
except ImportError:
    fitz = None
try:
    import pdfplumber
except ImportError:
    pdfplumber = None

DOC_EXT = {".pdf": "pdf", ".xlsx": "xlsx", ".xlsm": "xlsx", ".xls": "xls", ".csv": "csv", ".ics": "ics", ".ical": "ics", ".ifb": "ics"}
SKIP_EXT = (".jpg", ".jpeg", ".png", ".gif", ".bmp", ".svg", ".webp", ".ico", ".tif", ".tiff", ".zip", ".lzh", ".gz", ".7z",
            ".exe", ".msi", ".dmg", ".mp3", ".mp4", ".m4a", ".wav", ".mov", ".avi", ".wmv", ".flv", ".doc", ".docx", ".ppt",
            ".pptx", ".jtd", ".odt", ".ods", ".txt", ".js", ".css", ".xml", ".json", ".rss", ".rdf", ".atom", ".kml", ".kmz")
EXCERPT = 2000                     # 抜粋の本文の長さ
EXCERPT_MAX = EXCERPT + 1200       # 題名・見出し・表の中身を足した抜粋の上限
MAIN_MARK = ""               # 本文の始め (<main> など) の印。読んだあとで消す
CONTENT_IDS = {"main", "contents", "content", "main_body", "mainbody", "main-body", "main_contents", "main-contents",
               "maincontents", "main_content", "main-content", "maincontent", "tmp_contents", "page-content", "page_content",
               "pagebody", "honbun", "article-body", "primary"}



def _join(base, href):
    """urljoin。壊れたリンク (全角のコロンなど) で ValueError が出たら None"""
    try:
        return urllib.parse.urljoin(base, href)
    except ValueError:
        return None

def _err(e):
    """誤りの文 (手元の場所は消す)"""
    return scrub(f"{type(e).__name__}: {e}")[:200]


def ext_of(url):
    path = urllib.parse.unquote(urllib.parse.urlsplit(url).path).lower()
    m = re.search(r"\.[a-z0-9]{1,5}$", path)
    return m.group(0) if m else ""


def file_type(url, content_type="", head=b""):
    """html / pdf / xlsx / xls / csv / ics / other"""
    ct = (content_type or "").lower()
    ext = ext_of(url)
    h = head[:512].lstrip().lower()
    looks_html = h.startswith((b"<!doctype html", b"<html", b"<head")) or (h.startswith(b"<") and b"<html" in head[:4096].lower())
    if head.startswith(b"%PDF") or ("pdf" in ct and not looks_html):
        return "pdf"
    if "pdf" in ct or (ext == ".pdf" and looks_html):
        return "html"                                     # PDF のはずが HTML (見つからないページ・移転の案内など)
    if "text/calendar" in ct or head.lstrip().startswith(b"BEGIN:VCALENDAR"):
        return "ics"
    if "spreadsheetml" in ct or (head.startswith(b"PK\x03\x04") and ext in (".xlsx", ".xlsm")):
        return "xlsx"
    if "csv" in ct or ext == ".csv":
        return "csv"
    if "ms-excel" in ct or ext == ".xls" or head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    if "html" in ct or "xhtml" in ct:
        return "html"
    if ext in DOC_EXT:
        return DOC_EXT[ext]
    if h.startswith((b"<!doctype html", b"<html", b"<head", b"<meta", b"<!--")) or (not ct and b"<html" in head[:4096].lower()):
        return "html"
    return "other"


def nfkc(s):
    return unicodedata.normalize("NFKC", s or "")


# ---- 文字コード ----

_CHARSET_ALIAS = {"shift_jis": "cp932", "shift-jis": "cp932", "sjis": "cp932", "x-sjis": "cp932", "windows-31j": "cp932",
                  "ms_kanji": "cp932", "euc-jp": "euc_jp", "x-euc-jp": "euc_jp", "eucjp": "euc_jp", "eucjp-ms": "euc_jp",
                  "iso-2022-jp": "iso2022_jp", "utf8": "utf-8"}
# 宣言された文字コードで読めないときに試す、広い文字コード (EUC-JP の①② (NEC の特殊文字) は euc_jis_2004 で読める)
_WIDER = {"euc_jp": ("euc_jis_2004",), "cp932": ("shift_jis_2004",), "iso2022_jp": ("iso2022_jp_2004", "iso2022_jp_ext")}


def _charset(name):
    name = (name or "").strip().strip("\"'").lower()
    return _CHARSET_ALIAS.get(name, name) or None


def decode(body, content_type=""):
    """HTML などの bytes → (文字列, 使った文字コード)。ヘッダー → <meta> (と、その広い文字コード) → utf-8 → cp932 → euc_jp の順に
    厳しく読み、どれも読めなければ (はぐれた 1 バイトなど)、置き換え文字のいちばん少ないもので読む (宣言されたものを先に)"""
    declared = []
    m = re.search(r"charset=([\w\-]+)", content_type or "", re.I)
    if m:
        declared.append(_charset(m.group(1)))
    m = re.search(rb"<meta[^>]+charset=[\"']?([\w\-]+)", body[:4096], re.I)
    if m:
        declared.append(_charset(m.group(1).decode("ascii", "ignore")))
    cands = ["utf-8-sig"] if body.startswith(b"\xef\xbb\xbf") else []
    for cs in declared:
        cands += [cs, *_WIDER.get(cs, ())]
    cands += ["utf-8", "cp932", "euc_jp"] + ([] if declared else ["euc_jis_2004"])
    tried = []
    for cs in cands:
        if not cs or cs in tried:
            continue
        try:
            t = body.decode(cs)
        except LookupError:
            continue
        except UnicodeDecodeError:
            tried.append(cs)
            continue
        if cs == "cp932" and cs not in declared and _halfkana_heavy(t):
            tried.append(cs)                              # EUC-JP を cp932 で読むと半角カナだらけになる。次を試す
            continue
        return t, cs
    best = None
    for cs in tried:
        t = body.decode(cs, "replace")
        bad = t.count("�") + (len(t) if cs == "cp932" and cs not in declared and _halfkana_heavy(t) else 0)
        if best is None or bad < best[0]:
            best = (bad, t, cs)
    return (best[1], best[2]) if best else (body.decode("utf-8", "replace"), "utf-8")


def _halfkana_heavy(t):
    wide = [c for c in t[:20000] if ord(c) > 0x7f]
    return len(wide) >= 20 and sum(1 for c in wide if "｡" <= c <= "ﾟ") / len(wide) > 0.3


# ---- HTML ----

_BLOCK = {"p", "div", "br", "li", "tr", "td", "th", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "dd", "dt",
          "table", "ul", "ol", "header", "footer", "nav", "main", "aside", "caption", "option"}
_SKIP = {"script", "style", "noscript", "template", "svg"}


class _Parser(html.parser.HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.title, self.in_title = [], False
        self.heads, self.head = [], None
        self.links, self.a = [], None
        self.text, self.skip = [], 0
        self.tables, self.done_tables = [], []
        self.forms, self.form = [], None
        self.base = self.refresh = None
        self.nofollow = False
        self.meta_desc = ""
        self.main_marked = False
        self.canonical = None
        self.scripts = []

    def handle_starttag(self, tag, attrs):
        a = {k: (v or "") for k, v in attrs}
        if tag == "script" and a.get("src") and len(self.scripts) < 60:
            self.scripts.append(a["src"])                 # 埋め込みのアプリ (ごみスケなど) を見つけるため
        if tag in _SKIP:
            self.skip += 1
            return
        if not self.main_marked and not self.skip and (tag == "main" or a.get("role", "").lower() == "main" or (
                tag in ("div", "section", "article") and a.get("id", "").lower() in CONTENT_IDS)):
            self.text.append(MAIN_MARK)                   # 本文の始め (抜粋はここから)
            self.main_marked = True
        if tag == "title":
            self.in_title = True
        elif tag in ("h1", "h2", "h3"):
            self.head = [tag, []]
        elif tag == "a":
            self._close_a()
            if a.get("href"):
                self.a = [a["href"], [], a.get("title", "")]
        elif tag == "img" and self.a is not None and a.get("alt"):
            self.a[1].append(" " + a["alt"] + " ")
        elif tag == "area" and a.get("href"):
            self.links.append((a["href"], a.get("alt", ""), "area"))
        elif tag in ("iframe", "frame") and a.get("src"):
            self.links.append((a["src"], a.get("title", ""), tag))
        elif tag == "base" and a.get("href") and not self.base:
            self.base = a["href"]
        elif tag == "meta":
            name, he = a.get("name", "").lower(), a.get("http-equiv", "").lower()
            if he == "refresh":
                m = re.search(r"url\s*=\s*['\"]?([^'\"\s>]+)", a.get("content", ""), re.I)
                if m:
                    self.refresh = m.group(1)
            elif name == "robots" and "nofollow" in a.get("content", "").lower():
                self.nofollow = True
            elif name == "description":
                self.meta_desc = a.get("content", "")[:300]
            elif a.get("property", "").lower() == "og:url" and a.get("content") and not self.canonical:
                self.canonical = a["content"]
        elif tag == "link" and "canonical" in a.get("rel", "").lower().split() and a.get("href"):
            self.canonical = a["href"]                    # <link rel="canonical"> を og:url より先に
        elif tag == "table":
            self.tables.append({"rows": 0, "cells": [], "cur": None, "grid": []})
        elif tag == "tr" and self.tables:
            self._flush_cell()
            self.tables[-1]["rows"] += 1
            if len(self.tables[-1]["grid"]) < 300:
                self.tables[-1]["grid"].append([])
        elif tag in ("td", "th") and self.tables:
            self._flush_cell()
            self.tables[-1]["cur"] = []
        elif tag == "form":
            self.form = {"method": a.get("method", "get").lower() or "get", "action": a.get("action", ""), "inputs": [],
                         "role": a.get("role", "").lower()}
        elif tag == "input" and self.form is not None:
            self.form["inputs"].append(((a.get("type") or "text").lower(), a.get("name", ""), a.get("value", "")))
        if tag in _BLOCK:
            self.text.append("\n")

    def handle_endtag(self, tag):
        if tag in _SKIP:
            self.skip = max(0, self.skip - 1)
            return
        if tag == "title":
            self.in_title = False
        elif tag in ("h1", "h2", "h3") and self.head:
            t = " ".join("".join(self.head[1]).split())
            if t:
                self.heads.append((self.head[0], t[:200]))
            self.head = None
        elif tag == "a":
            self._close_a()
        elif tag in ("td", "th"):
            self._flush_cell()
        elif tag == "table" and self.tables:
            self._flush_cell()
            self.done_tables.append(self.tables.pop())
        elif tag == "form" and self.form is not None:
            self.forms.append(self.form)
            self.form = None
        if tag in _BLOCK:
            self.text.append("\n")

    def handle_data(self, data):
        if self.in_title:
            self.title.append(data)
            return
        if self.skip:
            return
        self.text.append(data)
        if self.head:
            self.head[1].append(data)
        if self.a is not None:
            self.a[1].append(data)
        if self.tables and self.tables[-1]["cur"] is not None:
            self.tables[-1]["cur"].append(data)

    def _close_a(self):
        if self.a is not None:
            text = " ".join("".join(self.a[1]).split()) or self.a[2]
            self.links.append((self.a[0], text[:200], "a"))
            self.a = None

    def _flush_cell(self):
        if self.tables and self.tables[-1]["cur"] is not None:
            t = self.tables[-1]
            if len(t["cells"]) < 3000:
                cell = " ".join("".join(t["cur"]).split())
                t["cells"].append(cell)
                if t["grid"] and t["rows"] == len(t["grid"]):
                    t["grid"][-1].append(cell)
            t["cur"] = None

    def finish(self):
        self._close_a()
        while self.tables:
            self._flush_cell()
            self.done_tables.append(self.tables.pop())
        if self.form is not None:
            self.forms.append(self.form)


class Page:
    """HTML を読んだもの。main_start は本文 (<main>・id="contents" など) の始めの位置 (無ければ None)。
    canonical は <link rel="canonical"> か og:url (絶対の URL にしたもの)、scripts は <script src> の URL"""

    def __init__(self, url, title, headings, text, links, tables, forms, base, refresh, nofollow, charset, description="",
                 main_start=None, canonical=None, scripts=()):
        self.url, self.title, self.headings, self.text = url, title, headings, text
        self.links, self.tables, self.forms = links, tables, forms
        self.base, self.refresh, self.nofollow, self.charset = base, refresh, nofollow, charset
        self.description, self.main_start = description, main_start
        self.canonical, self.scripts = canonical, list(scripts)


def parse_html(body, url, content_type=""):
    text, cs = decode(body, content_type)
    text = text.replace(MAIN_MARK, "")
    p = _Parser()
    try:
        p.feed(text)
        p.close()
    except Exception:                                  # 壊れた HTML でも読めたところまで使う
        pass
    p.finish()
    body_text = re.sub(r"[ \t\r\f\v　]+", " ", "".join(p.text))
    body_text = re.sub(r"\s*\n\s*", "\n", body_text).strip()
    main_start = body_text.find(MAIN_MARK)
    if main_start >= 0:
        body_text = body_text.replace(MAIN_MARK, "")
    base = (_join(url, p.base) or url) if p.base else url
    if not p.refresh and len(p.links) < 3:                 # リンクがほとんど無いトップの JavaScript の転送 (location.href = "…")
        m = re.search(r"location(?:\.href)?\s*=\s*[\"']([^\"'\s]+)[\"']", text[:20000])
        if m:
            p.refresh = m.group(1)
    canonical = _join(base, p.canonical.strip()) if p.canonical and p.canonical.strip() else None
    scripts = [u for u in (_join(base, s.strip()) for s in p.scripts if s.strip()) if u]
    return Page(url, " ".join("".join(p.title).split())[:300], p.heads[:30], body_text, p.links, p.done_tables, p.forms,
                base, p.refresh, p.nofollow, cs, p.meta_desc, main_start if main_start >= 0 else None, canonical, scripts)


def content_start(page):
    """本文の始め: <main> などの印 → 題名に含まれる h1 のうち本文でいちばん後に出るもの (ヘッダーのサイト名の h1 より、
    ページ名の h1 を選ぶ) → 0。ヘッダーやメニューの文字を抜粋に入れないため"""
    if page.main_start is not None:
        return page.main_start
    pos = [page.text.find(h) for tag, h in page.headings if tag == "h1" and len(h) >= 2 and h in page.title]
    pos = [i for i in pos if i >= 0]
    return max(pos) if pos else 0


def table_dump(tables, limit=800):
    """表の中身を短い文字に (1 行 = 1 行、セルは | で区切る)"""
    out, n = [], 0
    for t in tables:
        for row in t.get("grid") or []:
            line = " | ".join(c for c in row if c)
            if not line:
                continue
            out.append(line[:200])
            n += len(line) + 1
            if n >= limit:
                return "\n".join(out)[:limit]
        out.append("")
    return "\n".join(out).strip()[:limit]


def excerpt(page, tables=()):
    """LLM に渡す抜粋: 題名・見出し・収集日らしい表の中身・本文の始め (<main> か題名の h1 から 2,000 字)"""
    i = content_start(page)
    parts = [page.title] + [h for _, h in page.headings][:8]
    dump = table_dump(tables)
    if dump:
        parts.append("[表]\n" + dump)
    parts.append(page.text[i:i + EXCERPT])
    return "\n".join(p for p in parts if p)[:EXCERPT_MAX]


SEARCH_ENGINES = ("google.", "bing.com", "yahoo.co.jp", "yahoo.com", "duckduckgo.com", "baidu.com", "goo.ne.jp", "naver.")


def find_search_form(page):
    """サイト内検索が素直な GET のフォームなら (action, [(名前, 値)], 文字を入れる欄の名前) を返す"""
    for f in page.forms:
        if f["method"] != "get":
            continue
        texts = [(n, v) for t, n, v in f["inputs"] if t in ("text", "search", "") and n]
        if len(texts) != 1:
            continue
        name = texts[0][0]
        hidden = [(n, v) for t, n, v in f["inputs"] if t == "hidden" and n]
        if any(n.lower() in ("cx", "cof") for n, _ in hidden):        # Google のサイト内検索 (検索エンジン) は使わない
            continue
        if re.search(r"mail|name|tel|zip|pass", name, re.I):
            continue
        action = _join(page.base, f["action"] or page.url)
        if not action:
            continue
        if any(s in (urllib.parse.urlsplit(action).hostname or "") for s in SEARCH_ENGINES):
            continue
        return action, hidden, name
    return None


# ---- PDF ----

def _meaningful(s):
    return sum(1 for c in s if c.isalnum() or "぀" <= c <= "ヿ" or "一" <= c <= "鿿")


def _broken(s):
    """文字はあるが読めない (私用領域・置き換え文字・(cid:) だらけ、または漢字ばかりで仮名がほぼ無い)"""
    chars = [c for c in s if not c.isspace()]
    if not chars:
        return False
    bad = sum(1 for c in chars if c == "�" or "" <= c <= "" or ord(c) < 32) + 6 * s.count("(cid:")
    if bad / len(chars) > 0.3:
        return True
    kana = sum(1 for c in chars if "぀" <= c <= "ヿ")
    kanji = sum(1 for c in chars if "一" <= c <= "鿿")
    return kanji >= 100 and kana / (kana + kanji) < 0.03


VECTOR_CONTENT_BYTES = 1_000_000   # 1 ページ目の描画の命令がこれより大きければ、数えずに「輪郭線で描いた PDF」とする (遅いので)
PDF_MAX_PAGES = 80                 # 文字を読むページの上限 (PyMuPDF。pdfplumber は遅いので 30)
PDF_MAX_PAGES_PLUMBER = 30
PDF_MAX_CHARS = 300_000            # 読む文字の上限 (全ページの合計。超えたページは空とする)
PAGE_SEP = "\f"                    # analyze_pdf が返す文字のページの区切り
IMAGE_PAGE_CHARS = 20              # 意味のある文字がこれより少なく、画像のあるページは「画像だけのページ」


def _read(path, data):
    return pathlib.Path(path).read_bytes() if data is None else data


def split_pages(text):
    """analyze_pdf の文字 → ページごとの文字のリスト"""
    return text.split(PAGE_SEP) if text else []


def analyze_pdf(path=None, data=None):
    """PDF → 特徴 (pages, text_chars, images, drawings, image_only_pages, pdf_kind = text / image / vector / broken_text /
    empty / unreadable / not_pdf / unknown) と 1 ページ目の文字と、全ページの文字 (80 ページ・30 万字まで。
    ページは PAGE_SEP で区切る)。画像だけのページ (文字がほぼ無く画像がある) は数と始めのページ番号を記録する。
    ファイルは bytes で読んでから開く (PyMuPDF の誤りの文にはファイルの場所が入るため)"""
    feats = {"pdf_kind": "unknown"}
    first, text = "", ""
    try:
        data = _read(path, data)
    except OSError as e:
        feats.update(pdf_kind="unreadable", error=_err(e))
        return feats, "", ""
    if not data:
        feats.update(pdf_kind="unreadable", error="中身が空 (0 バイト)")
        return feats, "", ""
    if b"%PDF" not in data[:1024]:
        feats.update(pdf_kind="not_pdf", error="PDF ではない (始めに %PDF が無い)")
        return feats, "", ""
    imgs = 0
    image_only = []                                       # 画像だけのページ (1 から数える)
    try:
        if fitz is not None:
            with fitz.open(stream=data, filetype="pdf") as doc:
                if doc.needs_pass:
                    feats.update(pdf_kind="unreadable", error="パスワードつきの PDF")
                    return feats, "", ""
                feats["pages"] = doc.page_count
                pages, total, read = [], 0, 0
                for i in range(min(PDF_MAX_PAGES, doc.page_count)):
                    pg = doc[i]
                    t = ""
                    if total < PDF_MAX_CHARS:
                        t = pg.get_text()
                        read += 1
                    total += len(t)
                    pages.append(t)
                    if _meaningful(t) < IMAGE_PAGE_CHARS and total < PDF_MAX_CHARS:
                        try:
                            if pg.get_images():
                                image_only.append(i + 1)
                        except Exception:
                            pass
                meta_title = (doc.metadata or {}).get("title") or ""
                if doc.page_count:
                    p0 = doc[0]
                    imgs = feats["images_first_page"] = len(p0.get_images())
                    if _meaningful("\n".join(pages[:3])) < 40 and not imgs:      # 輪郭線の数は文字も画像も無いときだけ数える
                        try:
                            if len(p0.read_contents()) > VECTOR_CONTENT_BYTES:
                                feats["drawings_first_page"] = 100000
                            else:
                                feats["drawings_first_page"] = min(len(p0.get_drawings()), 100000)
                        except Exception:
                            pass
        elif pdfplumber is not None:
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                feats["pages"] = len(pdf.pages)
                pages, total, read = [], 0, 0
                for i, pg in enumerate(pdf.pages[:PDF_MAX_PAGES_PLUMBER]):
                    t = ""
                    if total < PDF_MAX_CHARS:
                        t = pg.extract_text() or ""
                        read += 1
                    total += len(t)
                    pages.append(t)
                    if _meaningful(t) < IMAGE_PAGE_CHARS and pg.images:
                        image_only.append(i + 1)
                if pdf.pages:
                    feats["images_first_page"] = len(pdf.pages[0].images)
                    feats["drawings_first_page"] = len(pdf.pages[0].curves) + len(pdf.pages[0].lines)
                meta_title = (pdf.metadata or {}).get("Title") or ""
        else:
            return feats, "", ""
    except Exception as e:
        feats["pdf_kind"] = "unreadable"
        feats["error"] = _err(e)
        return feats, "", ""
    if read < feats.get("pages", 0):
        feats["pages_read"] = read                        # 上限 (80 ページ・30 万字) で全部は読んでいない。読んだページの数
    if image_only:
        feats["image_only_pages"] = len(image_only)
        feats["image_only_first"] = image_only[0]
    first = pages[0] if pages else ""
    text = PAGE_SEP.join(pages)
    n = _meaningful(text)
    feats["text_chars"] = n
    if isinstance(meta_title, str) and meta_title.strip():
        feats["pdf_title"] = meta_title.strip()[:200]
    # 文字があるか: 始めの 3 ページで決める。3 ページに無くても、うしろに文字があり、画像だけのページが半分より少なければ文字の PDF
    # (画像だけのページが多い PDF の、うしろの 1 ページの短い文字では「文字の PDF」にしない)
    head = PAGE_SEP.join(pages[:3])
    has_text = _meaningful(head) >= 40 or (n >= 40 and 2 * len(image_only) < max(1, read))
    if has_text and not _broken(text):
        feats["pdf_kind"] = "text"
    elif has_text:
        feats["pdf_kind"] = "broken_text"
    elif feats.get("images_first_page") or image_only:
        feats["pdf_kind"] = "image"
    elif feats.get("drawings_first_page", 0) >= 50:
        feats["pdf_kind"] = "vector"
    else:
        feats["pdf_kind"] = "empty"
    return feats, first[:EXCERPT], text


def render_page(path, page_no=1, max_side=1600):
    """PDF の page_no ページ目 (1 から) を PNG (bytes) に。PyMuPDF が無い・ページが無ければ None。
    LLM の判定に添えるだけで、ファイルには残さない"""
    if fitz is None or not path:
        return None
    try:
        with fitz.open(stream=_read(path, None), filetype="pdf") as doc:
            if not 1 <= page_no <= doc.page_count:
                return None
            p = doc[page_no - 1]
            side = max(p.rect.width, p.rect.height) or 1
            zoom = min(2.0, max_side / side)
            return p.get_pixmap(matrix=fitz.Matrix(zoom, zoom)).tobytes("png")
    except Exception:
        return None


def render_first_page(path, max_side=1600):
    """PDF の 1 ページ目を PNG (bytes) に。PyMuPDF が無ければ None"""
    return render_page(path, 1, max_side)


def _squash(s):
    """空白をつめる (行の区切りは残す)"""
    s = re.sub(r"[ \t\r\f\v　]+", " ", s or "")
    return re.sub(r"\s*\n\s*", "\n", s).strip()


def pdf_excerpt(text, hits=(), best=(), head_chars=400, per_page=700, limit=EXCERPT_MAX, focus=None):
    """LLM に渡す PDF の抜粋。ごみの収集のページ (hits) があれば、ページの一覧と、1 ページ目の始め (head_chars 字) と、
    言葉の多いページ (best、多い順の 4 つ) の per_page 字 (focus の言葉が始めに出る少し前から) をページの順に並べる。
    無ければ前と同じ 1 ページ目の始め (前の判定の記録をそのまま使えるように)"""
    pts = split_pages(text)
    if not pts:
        return ""
    if not hits:
        return pts[0][:EXCERPT]
    hits = list(hits)
    shown = ", ".join(str(h) for h in hits[:30]) + (" ほか" if len(hits) > 30 else "")
    parts = [f"[PDF {len(pts)} ページ。ごみの収集の言葉の多いページ: {shown}]"]
    p1 = _squash(pts[0])
    if p1:
        parts.append("[p.1]\n" + p1[:head_chars])
    n = sum(len(x) + 1 for x in parts)
    pick = [b for b in (best or hits) if b != 1 and b <= len(pts)][:4]
    for no in sorted(pick):
        t = _squash(pts[no - 1])
        m = focus.search(t) if focus is not None else None
        start = max(0, m.start() - 80) if m else 0             # ページの始めの広告・行事の文字を飛ばす
        t = t[start:start + per_page]
        if not t or n + len(t) + 10 > limit:
            continue
        parts.append(f"[p.{no}]\n{t}")
        n += len(t) + 10
    return "\n".join(parts)[:limit]


# ---- CSV・Excel・iCalendar ----

def analyze_csv(body):
    text, cs = decode(body, "")
    rows = []
    try:
        for i, row in enumerate(csv.reader(io.StringIO(text))):
            rows.append(row)
            if i >= 2000:
                break
    except csv.Error:
        pass
    feats = {"rows": len(rows), "columns": max((len(r) for r in rows), default=0), "header": [c[:40] for c in (rows[0] if rows else [])][:30]}
    sample = "\n".join(",".join(r) for r in rows[:200])
    return feats, sample[:EXCERPT], sample


def analyze_xlsx(path=None, data=None):
    """Excel (xlsx) → 特徴 (シート名・読んだ行の数) と始めの中身。キャッシュのファイル名は .body なので、
    openpyxl には bytes で渡す (ファイル名で渡すと拡張子で断られる)"""
    try:
        import openpyxl
    except ImportError:
        return {"xlsx": "openpyxl が無いので見ていない"}, "", ""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(_read(path, data)), read_only=True, data_only=True)
    except Exception as e:
        return {"error": _err(e)}, "", ""
    lines, feats = [], {"sheets": wb.sheetnames[:20]}
    try:
        for ws in wb.worksheets[:3]:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 200:
                    break
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    lines.append(",".join(cells))
    except Exception as e:
        feats["error"] = _err(e)
    finally:
        wb.close()
    feats["rows_read"] = len(lines)
    sample = "\n".join(lines)
    return feats, sample[:EXCERPT], sample


def analyze_ics(body):
    text, _ = decode(body, "")
    text = re.sub(r"\r?\n[ \t]", "", text)              # 折り返しを戻す
    events = len(re.findall(r"^BEGIN:VEVENT", text, re.M))
    summaries = re.findall(r"^SUMMARY[^:]*:(.*)$", text, re.M)
    starts = sorted(re.findall(r"^DTSTART[^:]*:(\d{8})", text, re.M))
    feats = {"events": events, "first_date": starts[0] if starts else None, "last_date": starts[-1] if starts else None,
             "summaries": sorted(set(s.strip()[:40] for s in summaries))[:20]}
    sample = "\n".join(summaries[:200])
    return feats, sample[:EXCERPT], sample
