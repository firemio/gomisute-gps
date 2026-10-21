"""入力の準備: 資料 1 本 → data/llmread/inputs/<doc_id>/ (2 つのモデルに同じものを渡す)

  HTML      → 本文の範囲 (htmlbody.py。ページの形だけで決める) を文字に。表は rowspan・colspan だけ残した簡単な HTML の表にする
  CSV・Excel → 表をそのまま markdown の表に (Excel はシートごと。結合したマスは左上の値で埋める)
  PDF       → ページごとの PNG (既定 150dpi) と、文字層があればその文字

意味の読み取りはしない (どの行が地区か・どの文が祝日の扱いかは決めない)。決めるのは LLM。
入力はキャッシュ (data/discover/cache/) と手元のファイル (data/kanazawa_cal/ など) だけ。サイトには取りに行かない。
"""
import csv, html as htmllib, io, re, unicodedata

from bs4 import Comment, NavigableString, Tag

from common import INPUTS, ROOT, now_utc, sha256, write_json, read_json
import decode, fetch_docs, fetcher, htmlbody                     # noqa: E402  (fetcher は nationwide/discover)

HEADINGS = {f"h{i}": i for i in range(1, 7)}
BLOCKS = {"p", "div", "li", "dt", "dd", "ul", "ol", "dl", "section", "article", "blockquote", "pre", "caption", "figure",
          "figcaption", "header", "footer", "main", "nav", "aside", "form", "address", "tr", "td", "th", "tbody", "thead"}
DPI = 150


def _ws(s):
    return re.sub(r"[ \t\r\n　\xa0]+", " ", s).strip()


# ---- 取り出し ----
def load(doc):
    """計画の 1 件 → (bytes, 記録)。url はキャッシュから、path は手元のファイルから"""
    if doc.get("path"):
        p = ROOT / doc["path"]
        return p.read_bytes(), {"path": doc["path"], "content_type": doc.get("content_type") or ""}
    cache = fetcher.Cache(fetch_docs.CACHE_DIR)
    body, meta = fetch_docs.body_of(cache, fetcher.norm_url(doc["url"]) or doc["url"])
    if body is None:
        raise FileNotFoundError(f"キャッシュに無い: {doc['url']}")
    hdr = (meta or {}).get("headers") or {}
    return body, {"url": doc["url"], "content_type": hdr.get("content-type") or "", "fetched_at": (meta or {}).get("fetched_at"),
                  "last_modified": hdr.get("last-modified")}


def kind_of(body, url="", content_type=""):
    if body[:5] == b"%PDF-":
        return "pdf"
    if body[:4] == b"PK\x03\x04":
        return "xlsx"
    if body[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    u = url.lower().split("?")[0]
    if u.endswith(".csv") or "csv" in content_type.lower():
        return "csv"
    return "html"


# ---- HTML ----
def _cell_text(c):
    for br in c.find_all("br"):
        br.replace_with(" ")
    return _ws(c.get_text(" "))


def table_html(t):
    """表 → 簡単な HTML (caption・tr・th・td と rowspan・colspan だけ)。入れ子の表は文字にする"""
    out = ["<table>"]
    cap = t.find("caption")
    if cap is not None and cap.find_parent("table") is t:
        out.append(f"<caption>{htmllib.escape(_cell_text(cap), quote=False)}</caption>")
    for tr in t.find_all("tr"):
        if tr.find_parent("table") is not t:
            continue
        cells = []
        for c in tr.find_all(["th", "td"], recursive=False):
            attrs = "".join(f' {a}="{c.get(a)}"' for a in ("rowspan", "colspan") if c.get(a) and str(c.get(a)) != "1")
            cells.append(f"<{c.name}{attrs}>{htmllib.escape(_cell_text(c), quote=False)}</{c.name}>")
        out.append("<tr>" + "".join(cells) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def render(root):
    """本文の要素 → 文字 (見出しは #、表は簡単な HTML、箇条は -)"""
    lines, buf = [], []

    def flush():
        t = _ws("".join(buf))
        if t and t != "-":
            lines.append(t)
        buf.clear()

    def walk(el):
        for ch in el.children:
            if isinstance(ch, Comment):
                continue
            if isinstance(ch, NavigableString):
                buf.append(str(ch))
                continue
            if not isinstance(ch, Tag) or ch.name in ("script", "style", "noscript", "template"):
                continue
            if ch.name == "table":
                flush()
                lines.append(table_html(ch))
            elif ch.name in HEADINGS:
                flush()
                t = _ws(ch.get_text(" "))
                if t:
                    lines.append("#" * HEADINGS[ch.name] + " " + t)
            elif ch.name == "br":
                flush()
            elif ch.name == "img":
                alt = _ws(ch.get("alt") or "")
                if alt:
                    buf.append(f"[画像: {alt}]")
            elif ch.name == "li":
                flush()
                buf.append("- ")
                walk(ch)
                flush()
            elif ch.name in BLOCKS:
                flush()
                walk(ch)
                flush()
            else:
                walk(ch)
    walk(root)
    flush()
    return "\n".join(lines)


def html_text(body, content_type="", url=""):
    """HTML の bytes → (本文の文字, 記録)"""
    d = decode.decode(body, content_type)
    host = fetcher.norm_url(url).split("/")[2] if url.startswith("http") else None
    bp = fetch_docs.load_boilerplate(host) if host else None
    region = htmlbody.content_region(d.text, boilerplate=bp)
    title = None
    try:
        soup = htmlbody.parse(d.text)
        title = _ws(soup.title.get_text()) if soup.title else None
    except Exception:
        pass
    text = render(region.root)
    return text, {"encoding": d.encoding, "region": region.method, "title": title}


def _nk(s):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", s))


def section(text, heading):
    """本文の文字から、heading を含む見出しの行から、同じかより上の見出しの手前までを切り出す (無ければ None)"""
    ls = text.split("\n")
    for i, ln in enumerate(ls):
        m = re.match(r"^(#+) ", ln)
        if m and _nk(heading) in _nk(ln):
            lv = len(m.group(1))
            j = i + 1
            while j < len(ls) and not (re.match(r"^(#+) ", ls[j]) and len(re.match(r"^(#+) ", ls[j]).group(1)) <= lv):
                j += 1
            return "\n".join(ls[i:j])
    return None


# ---- CSV・Excel ----
def md_table(rows):
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    esc = lambda c: _ws(str(c)).replace("|", "｜")
    return "\n".join("| " + " | ".join(esc(c) for c in list(r) + [""] * (w - len(r))) + " |" for r in rows)


def read_csv(body, content_type=""):
    """CSV の bytes → (行のリスト, 文字コードの記録)"""
    d = decode.decode(body, content_type, html=False)
    text = d.text
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    return rows, d.record()


def xlsx_sheets(body):
    """Excel (xlsx) の bytes → [{"sheet", "hidden", "rows"}]。結合したマスは左上の値で埋める"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(body), data_only=True)
    out = []
    for ws in wb.worksheets:
        rows = [[("" if v is None else str(v)).strip() for v in r] for r in ws.iter_rows(values_only=True)]
        for rng in ws.merged_cells.ranges:
            v = rows[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(rows) else ""
            for r in range(rng.min_row - 1, min(rng.max_row, len(rows))):
                for c in range(rng.min_col - 1, min(rng.max_col, len(rows[r]))):
                    rows[r][c] = v
        out.append({"sheet": ws.title, "hidden": ws.sheet_state != "visible", "rows": rows})
    return out


def xls_sheets(body):
    """Excel (xls) の bytes → [{"sheet", "hidden", "rows"}]。結合したマスは左上の値で埋める"""
    import xlrd
    # 結合したマス (merged_cells) は formatting_info=True のときだけ読まれる (xlrd は False だと XL_MERGEDCELLS を飛ばす)
    wb = xlrd.open_workbook(file_contents=body, formatting_info=True)
    out = []
    for sh in wb.sheets():
        rows = [[str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)] for r in range(sh.nrows)]
        for (r0, r1, c0, c1) in getattr(sh, "merged_cells", []) or []:
            v = rows[r0][c0]
            for r in range(r0, r1):
                for c in range(c0, c1):
                    rows[r][c] = v
        out.append({"sheet": sh.name, "hidden": bool(getattr(sh, "visibility", 0)), "rows": rows})
    return out


def csv_text(body, content_type=""):
    rows, rec = read_csv(body, content_type)
    return md_table(rows), {"encoding": rec.get("encoding") if isinstance(rec, dict) else None, "rows": len(rows)}


def excel_text(body, kind):
    sheets = xlsx_sheets(body) if kind == "xlsx" else xls_sheets(body)
    parts = []
    for s in sheets:
        parts.append(f"## シート「{s['sheet']}」" + (" (隠しシート)" if s["hidden"] else ""))
        parts.append(md_table(s["rows"]))
    return "\n\n".join(parts), {"sheets": len(sheets)}


# ---- PDF ----
def pdf_pages(body, pages=None, dpi=DPI):
    """PDF → [(ページ番号 1 始まり, PNG の bytes, 文字層)]"""
    import fitz
    out = []
    with fitz.open(stream=body, filetype="pdf") as d:
        for n in (pages or range(1, len(d) + 1)):
            pg = d[n - 1]
            out.append((n, pg.get_pixmap(dpi=dpi).tobytes("png"), pg.get_text("text").strip()))
    return out


def pdf_page_count(body):
    import fitz
    with fitz.open(stream=body, filetype="pdf") as d:
        return len(d), sum(len(p.get_text("text").strip()) for p in d)


# ---- 1 件の準備 ----
def prepare(doc, dpi=DPI, force=False):
    """計画の 1 件 (doc_id・url か path・pages・section) → manifest (data/llmread/inputs/<doc_id>/manifest.json)"""
    d = INPUTS / doc["doc_id"]
    mf = d / "manifest.json"
    if mf.exists() and not force:
        return read_json(mf)
    body, rec = load(doc)
    kind = kind_of(body, doc.get("url") or doc.get("path") or "", rec.get("content_type") or "")
    d.mkdir(parents=True, exist_ok=True)
    for old in d.glob("page-*"):
        old.unlink()
    info, text, pages = {}, None, []
    if kind == "html":
        text, info = html_text(body, rec.get("content_type") or "", doc.get("url") or "")
        if doc.get("section"):
            sec = section(text, doc["section"])
            if sec is None:
                raise ValueError(f"見出し「{doc['section']}」が本文に無い: {doc['doc_id']}")
            text = sec
    elif kind == "csv":
        text, info = csv_text(body, rec.get("content_type") or "")
    elif kind in ("xlsx", "xls"):
        text, info = excel_text(body, kind)
    else:
        for n, png, t in pdf_pages(body, doc.get("pages"), dpi):
            name = f"page-{n:03d}"
            (d / f"{name}.png").write_bytes(png)
            if t:
                (d / f"{name}.txt").write_text(t, "utf-8")
            pages.append({"page": n, "png": f"{name}.png", "text": f"{name}.txt" if t else None, "chars": len(t),
                          "png_sha256": sha256(png), "png_bytes": len(png)})
    if text is not None:
        (d / "text.md").write_text(text, "utf-8")
    h = sha256((text or "") + "\0" + "\0".join(p["png_sha256"] for p in pages)
               + "\0" + "\0".join((d / p["text"]).read_text("utf-8") for p in pages if p["text"]))
    m = {"doc_id": doc["doc_id"], "target": doc.get("target"), "url": doc.get("url"), "path": doc.get("path"), "kind": kind,
         "source_sha256": sha256(body), "source_bytes": len(body), "fetched_at": rec.get("fetched_at"),
         "last_modified": rec.get("last_modified"), "section": doc.get("section"), "dpi": dpi if pages else None,
         "text": "text.md" if text is not None else None, "text_chars": len(text or "") + sum(p["chars"] for p in pages),
         "pages": pages, "info": info, "input_sha256": h, "prepared_at": now_utc().isoformat(timespec="seconds")}
    write_json(mf, m)
    return m
